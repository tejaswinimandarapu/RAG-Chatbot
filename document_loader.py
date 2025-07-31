# utils/document_loader.py
import docx, csv, pptx, PyPDF2
import os

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        with open(file_path, "rb") as f:
            return " ".join([page.extract_text() for page in PyPDF2.PdfReader(f).pages])
    elif ext == ".docx":
        return " ".join([para.text for para in docx.Document(file_path).paragraphs])
    elif ext == ".pptx":
        prs = pptx.Presentation(file_path)
        return " ".join([shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text")])
    elif ext == ".csv":
        with open(file_path, "r", encoding="utf-8") as f:
            return " ".join([" ".join(row) for row in csv.reader(f)])
    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    else:
        return ""
